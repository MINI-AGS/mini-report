from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

def crear_tabla(ws, type_table, table_name):
    # Definir el rango de la tabla en base al tamaño de los datos
    rango_tabla = f"A1:G{ws.max_row}"  

    # Crear la tabla
    tabla = Table(displayName=table_name, ref=rango_tabla)

    # Definir el estilo de la tabla
    estilo = TableStyleInfo(
        name=type_table,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )

    tabla.tableStyleInfo = estilo
    ws.add_table(tabla)